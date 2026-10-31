import io
import csv
from typing import List, Dict, Any, Tuple
import openpyxl
from pypdf import PdfReader
from app.schemas.schemas import MemberGazetteItem, GazettePreviewResponse

def normalize_text(text: Any) -> str:
    if text is None:
        return ""
    return str(text).strip()

def parse_gazette_spreadsheet(file_bytes: bytes, filename: str) -> Tuple[List[MemberGazetteItem], List[str]]:
    items: List[MemberGazetteItem] = []
    warnings: List[str] = []
    
    if filename.lower().endswith(".csv"):
        content = file_bytes.decode("utf-8", errors="ignore")
        reader = csv.DictReader(io.StringIO(content))
        for row in reader:
            normalized_row = {k.strip().lower(): v.strip() for k, v in row.items() if k}
            item = extract_row_data(normalized_row)
            if item:
                items.append(item)
    else:
        # Excel
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            
            # Find header row
            header_idx = -1
            headers = []
            for idx, r in enumerate(rows):
                str_r = [str(c).lower().strip() for c in r if c is not None]
                if any("surname" in x or "name" in x for x in str_r):
                    header_idx = idx
                    headers = [str(c).strip().lower() if c is not None else "" for c in r]
                    break
            
            if header_idx == -1:
                continue
            
            # Category inference from sheet name if applicable
            sheet_cat = None
            s_lower = sheetname.lower()
            if "woman" in s_lower or "cwr" in s_lower:
                sheet_cat = "WOMAN_REP"
            elif "nominated" in s_lower:
                sheet_cat = "NOMINATED"
            elif "mp" in s_lower or "constituency" in s_lower:
                sheet_cat = "MP"
            
            for row in rows[header_idx + 1:]:
                if not any(row):
                    continue
                row_dict = {}
                for h, val in zip(headers, row):
                    if h:
                        row_dict[h] = normalize_text(val)
                
                item = extract_row_data(row_dict, default_category=sheet_cat)
                if item:
                    items.append(item)
                    
    return items, warnings

def parse_gazette_pdf(file_bytes: bytes) -> Tuple[List[MemberGazetteItem], List[str]]:
    items: List[MemberGazetteItem] = []
    warnings: List[str] = []
    
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        current_cat = "MP"
        for page in reader.pages:
            text = page.extract_text() or ""
            lines = text.split("\n")
            for line in lines:
                l_str = line.strip()
                if "COUNTY WOMAN REPRESENTATIVE" in l_str.upper():
                    current_cat = "WOMAN_REP"
                    continue
                elif "NOMINATED" in l_str.upper():
                    current_cat = "NOMINATED"
                    continue
                elif "MEMBERS OF THE NATIONAL ASSEMBLY" in l_str.upper() or "CONSTITUENCY" in l_str.upper():
                    current_cat = "MP"
                    continue
                
                # Check for tabular row structure
                parts = [p.strip() for p in l_str.split("  ") if p.strip()]
                if len(parts) >= 3:
                    # heuristic
                    surname = parts[0]
                    other = parts[1]
                    party = parts[-1]
                    constituency = parts[2] if len(parts) >= 4 else None
                    county = parts[3] if len(parts) >= 5 else None
                    
                    full_name = f"{surname} {other}".strip()
                    items.append(MemberGazetteItem(
                        surname=surname,
                        other_names=other,
                        full_name=full_name,
                        category=current_cat,
                        constituency=constituency,
                        county=county,
                        party=party
                    ))
    except Exception as e:
        warnings.append(f"PDF parsing error: {str(e)}. Use Excel template if file is scanned.")
        
    return items, warnings

def extract_row_data(row: Dict[str, str], default_category: str = None) -> MemberGazetteItem:
    # Keys might be 'surname', 'other names', 'other_names', 'full name', 'full_name'
    surname = row.get("surname", "")
    other_names = row.get("other names") or row.get("other_names", "")
    full_name = row.get("full name") or row.get("full_name") or row.get("name", "")
    
    if not surname and not full_name:
        return None
        
    if not surname and full_name:
        parts = full_name.split()
        if len(parts) > 1:
            surname = parts[0]
            other_names = " ".join(parts[1:])
        else:
            surname = full_name
            other_names = ""
            
    if not full_name:
        full_name = f"{surname} {other_names}".strip()
        
    constituency = row.get("constituency") or row.get("constituency name") or None
    county = row.get("county") or row.get("county name") or None
    party = row.get("party") or row.get("political party") or "Independent"
    
    category = row.get("category") or default_category
    if not category:
        if constituency and county:
            category = "MP"
        elif county and not constituency:
            category = "WOMAN_REP"
        elif not constituency and not county:
            category = "NOMINATED"
        else:
            category = "MP"
            
    cat_upper = category.upper()
    if "WOMAN" in cat_upper or "CWR" in cat_upper:
        category = "WOMAN_REP"
    elif "NOMINAT" in cat_upper:
        category = "NOMINATED"
    elif "SPEAKER" in cat_upper:
        category = "SPEAKER"
    else:
        category = "MP"
        
    return MemberGazetteItem(
        surname=surname,
        other_names=other_names,
        full_name=full_name,
        category=category,
        constituency=constituency,
        county=county,
        party=party
    )

def preview_gazette(items: List[MemberGazetteItem], warnings: List[str]) -> GazettePreviewResponse:
    count_mp = sum(1 for x in items if x.category == "MP")
    count_cwr = sum(1 for x in items if x.category == "WOMAN_REP")
    count_nom = sum(1 for x in items if x.category == "NOMINATED")
    
    is_valid = (count_mp == 290 and count_cwr == 47 and count_nom == 12)
    
    if count_mp != 290:
        warnings.append(f"Elected Constituency MPs: found {count_mp}, expected 290.")
    if count_cwr != 47:
        warnings.append(f"County Woman Representatives: found {count_cwr}, expected 47.")
    if count_nom != 12:
        warnings.append(f"Nominated Members: found {count_nom}, expected 12.")
        
    return GazettePreviewResponse(
        total_parsed=len(items),
        count_mp=count_mp,
        count_woman_rep=count_cwr,
        count_nominated=count_nom,
        target_mp=290,
        target_woman_rep=47,
        target_nominated=12,
        is_counts_valid=is_valid,
        sample_records=items[:15],
        warnings=warnings
    )
